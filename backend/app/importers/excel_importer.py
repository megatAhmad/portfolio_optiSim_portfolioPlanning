"""
PlanningSpace Excel Template Parser

Imports portfolio data from Excel files following the PlanningSpace template format.
Uses openpyxl for reading .xlsx files.

Expected Excel Structure:
- Sheet "Opportunities": Project list with name, type, business_unit, country, etc.
- Sheet "Outcomes": Outcome definitions with opportunity_name, outcome_name, weight
- Sheet "Metrics": Time-series data with opportunity, outcome, metric_name, year columns
- Sheet "Attributes" (optional): Fixture/classification attributes per opportunity
- Sheet "Dependencies" (optional): Inter-project dependency definitions
- Sheet "Groups" (optional): Selection group definitions
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import structlog

from app.importers.validators import ImportValidator, ValidationResult

logger = structlog.get_logger(__name__)


class ExcelImportResult:
    """Container for parsed Excel data."""

    def __init__(self) -> None:
        self.opportunities: List[Dict[str, Any]] = []
        self.outcomes: Dict[str, List[Dict[str, Any]]] = {}
        self.metrics: Dict[str, Dict[str, Dict[str, List[float]]]] = {}
        self.attributes: Dict[str, Dict[str, Any]] = {}
        self.dependencies: List[Dict[str, Any]] = []
        self.groups: List[Dict[str, Any]] = []
        self.validation: Optional[ValidationResult] = None


class ExcelImporter:
    """Import portfolio data from PlanningSpace-compatible Excel files.

    Usage::

        importer = ExcelImporter(planning_horizon=30)
        result = importer.import_file("portfolio_data.xlsx")

        if result.validation.is_valid:
            # Persist to database
            ...
        else:
            # Report errors
            for error in result.validation.errors:
                print(f"{error.field}: {error.message}")
    """

    def __init__(self, planning_horizon: int = 30) -> None:
        self.planning_horizon = planning_horizon
        self.validator = ImportValidator(planning_horizon)

    def import_file(
        self,
        file_path: str | Path,
        validate: bool = True,
    ) -> ExcelImportResult:
        """Import all data from an Excel file.

        Args:
            file_path: Path to the .xlsx file
            validate: Run validation after parsing (default True)

        Returns:
            ExcelImportResult with all parsed data and validation results
        """
        try:
            import openpyxl
        except ImportError as e:
            raise ImportError(
                "openpyxl is required for Excel import. Install with: pip install openpyxl"
            ) from e

        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"Excel file not found: {file_path}")

        if not file_path.suffix.lower() in (".xlsx", ".xlsm"):
            raise ValueError(f"Unsupported file format: {file_path.suffix}. Expected .xlsx or .xlsm")

        logger.info("excel_import_started", file_path=str(file_path))

        wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
        result = ExcelImportResult()

        try:
            # Parse each sheet
            if "Opportunities" in wb.sheetnames:
                result.opportunities = self._parse_opportunities(wb["Opportunities"])
            elif "Projects" in wb.sheetnames:
                result.opportunities = self._parse_opportunities(wb["Projects"])

            if "Outcomes" in wb.sheetnames:
                result.outcomes = self._parse_outcomes(wb["Outcomes"])

            if "Metrics" in wb.sheetnames:
                result.metrics = self._parse_metrics(wb["Metrics"])

            if "Attributes" in wb.sheetnames:
                result.attributes = self._parse_attributes(wb["Attributes"])

            if "Dependencies" in wb.sheetnames:
                result.dependencies = self._parse_dependencies(wb["Dependencies"])

            if "Groups" in wb.sheetnames:
                result.groups = self._parse_groups(wb["Groups"])

        finally:
            wb.close()

        # Validate if requested
        if validate:
            result.validation = self.validator.validate_full_import(
                result.opportunities, result.outcomes, result.metrics
            )

        logger.info(
            "excel_import_completed",
            opportunities=len(result.opportunities),
            outcome_groups=len(result.outcomes),
            metric_types=len(result.metrics),
            is_valid=result.validation.is_valid if result.validation else None,
        )

        return result

    def _parse_opportunities(self, sheet: Any) -> List[Dict[str, Any]]:
        """Parse opportunities from the Opportunities/Projects sheet."""
        opportunities: List[Dict[str, Any]] = []

        # Read header row
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        header_map = {str(h).lower().strip(): i for i, h in enumerate(headers) if h}

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue

            opp: Dict[str, Any] = {}

            # Map columns by header name
            name_idx = header_map.get("name", header_map.get("opportunity", 0))
            opp["name"] = str(row[name_idx]).strip() if row[name_idx] else ""

            for key in ("type", "business_unit", "country", "location_lat", "location_lon", "description"):
                if key in header_map and row[header_map[key]] is not None:
                    opp[key] = row[header_map[key]]

            if opp["name"]:
                opportunities.append(opp)

        return opportunities

    def _parse_outcomes(self, sheet: Any) -> Dict[str, List[Dict[str, Any]]]:
        """Parse outcomes from the Outcomes sheet."""
        outcomes: Dict[str, List[Dict[str, Any]]] = {}

        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        header_map = {str(h).lower().strip(): i for i, h in enumerate(headers) if h}

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue

            opp_idx = header_map.get("opportunity", header_map.get("opportunity_name", 0))
            name_idx = header_map.get("outcome", header_map.get("outcome_name", header_map.get("name", 1)))
            weight_idx = header_map.get("weight", header_map.get("probability", 2))

            opp_name = str(row[opp_idx]).strip() if row[opp_idx] else ""
            outcome_name = str(row[name_idx]).strip() if row[name_idx] else ""
            weight = float(row[weight_idx]) if row[weight_idx] is not None else 0.0

            if opp_name and outcome_name:
                if opp_name not in outcomes:
                    outcomes[opp_name] = []
                outcomes[opp_name].append({
                    "name": outcome_name,
                    "weight": weight,
                })

        return outcomes

    def _parse_metrics(
        self,
        sheet: Any,
    ) -> Dict[str, Dict[str, Dict[str, List[float]]]]:
        """Parse time-series metrics from the Metrics sheet.

        Expected columns: opportunity, outcome, metric_name, year_0, year_1, ..., year_N
        """
        metrics: Dict[str, Dict[str, Dict[str, List[float]]]] = {}

        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        header_map = {str(h).lower().strip(): i for i, h in enumerate(headers) if h}

        # Identify year columns (columns that start with "year_" or are numeric)
        year_columns: List[int] = []
        for i, h in enumerate(headers):
            if h is None:
                continue
            h_str = str(h).lower().strip()
            if h_str.startswith("year_") or h_str.isdigit():
                year_columns.append(i)

        # If no explicit year columns, assume columns after the first 3 are year data
        if not year_columns and len(headers) > 3:
            year_columns = list(range(3, len(headers)))

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue

            opp_idx = header_map.get("opportunity", header_map.get("opportunity_name", 0))
            out_idx = header_map.get("outcome", header_map.get("outcome_name", 1))
            metric_idx = header_map.get("metric", header_map.get("metric_name", 2))

            opp_name = str(row[opp_idx]).strip() if row[opp_idx] else ""
            outcome = str(row[out_idx]).strip() if row[out_idx] else ""
            metric_name = str(row[metric_idx]).strip() if row[metric_idx] else ""

            if not (opp_name and outcome and metric_name):
                continue

            # Extract year values
            values: List[float] = []
            for col_idx in year_columns:
                if col_idx < len(row) and row[col_idx] is not None:
                    try:
                        values.append(float(row[col_idx]))
                    except (ValueError, TypeError):
                        values.append(0.0)
                else:
                    values.append(0.0)

            # Store in nested dict
            if metric_name not in metrics:
                metrics[metric_name] = {}
            if opp_name not in metrics[metric_name]:
                metrics[metric_name][opp_name] = {}
            metrics[metric_name][opp_name][outcome] = values

        return metrics

    def _parse_attributes(self, sheet: Any) -> Dict[str, Dict[str, Any]]:
        """Parse opportunity attributes from the Attributes sheet."""
        attributes: Dict[str, Dict[str, Any]] = {}

        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        header_map = {str(h).lower().strip(): i for i, h in enumerate(headers) if h}

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue

            opp_idx = header_map.get("opportunity", header_map.get("name", 0))
            opp_name = str(row[opp_idx]).strip() if row[opp_idx] else ""

            if not opp_name:
                continue

            attrs: Dict[str, Any] = {}
            for key, idx in header_map.items():
                if key not in ("opportunity", "name") and idx < len(row):
                    attrs[key] = row[idx]

            attributes[opp_name] = attrs

        return attributes

    def _parse_dependencies(self, sheet: Any) -> List[Dict[str, Any]]:
        """Parse dependency definitions from the Dependencies sheet."""
        dependencies: List[Dict[str, Any]] = []

        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        header_map = {str(h).lower().strip(): i for i, h in enumerate(headers) if h}

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue

            dep: Dict[str, Any] = {}
            for key, idx in header_map.items():
                if idx < len(row):
                    dep[key] = row[idx]

            dependencies.append(dep)

        return dependencies

    def _parse_groups(self, sheet: Any) -> List[Dict[str, Any]]:
        """Parse selection group definitions from the Groups sheet."""
        groups: List[Dict[str, Any]] = []

        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        header_map = {str(h).lower().strip(): i for i, h in enumerate(headers) if h}

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue

            group: Dict[str, Any] = {}
            for key, idx in header_map.items():
                if idx < len(row):
                    group[key] = row[idx]

            groups.append(group)

        return groups
